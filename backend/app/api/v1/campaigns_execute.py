from fastapi import APIRouter, Depends, HTTPException, Header, UploadFile, File, Form
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import Optional
from app.core.database import get_db
from app.core.security import decode_token
from app.core.config import settings as app_settings
from app.models import Execution, EmailLog, SMTPProfile, Mapping, MappingEntry, Notification, Setting, User
import os
import re
import tempfile
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.image import MIMEImage
from email import encoders
from datetime import datetime
import base64
import pandas as pd
import openpyxl
from app.services.excel_service import (
    detect_active_range, render_excel_range_html, render_excel_range_image,
    _plain_table_html, prewarm_xlsb_conversions, prewarm_summary_images,
    get_cell_value,
)
from app.services.mascot_service import generate_bounce_gif

router = APIRouter(prefix="/campaigns", tags=["campaigns"])

# {{Cell:B3}} resolves against the campaign's configured sheet; {{Cell:Sheet!B3}}
# names its own sheet. Kept in sync with frontend/components/campaigns/cell-reference-field.tsx.
_CELL_PLACEHOLDER_RE = re.compile(r"\{\{Cell:(?:([^!{}]+)!)?([A-Za-z]{1,3}[0-9]+)\}\}")


def _resolve_cell_placeholders(text: str, file_path: str, default_sheet_name: str) -> str:
    if "{{Cell:" not in text:
        return text

    def _replace(match: "re.Match") -> str:
        if not file_path or not os.path.exists(file_path):
            return ""
        sheet = (match.group(1) or default_sheet_name or "Summary").strip()
        cell_ref = match.group(2).upper()
        return get_cell_value(file_path, sheet, cell_ref)

    return _CELL_PLACEHOLDER_RE.sub(_replace, text)

class CampaignExecuteRequest(BaseModel):
    smtp_profile: str
    mapping_id: Optional[int] = None
    subject: str
    body_template: str
    report_type: str
    sheet_name: str = "Summary"
    start_cell: str = ""  # empty = auto-detect the active range
    summary_format: str = "table"  # "table" (HTML) or "image" (LibreOffice PNG)
    attach_file: bool = True
    campaign_folder: str

def get_current_user(authorization: str = Header(...)):
    token = authorization.replace("Bearer ", "")
    payload = decode_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="Invalid token")
    return int(payload.get("sub")), payload.get("role", "viewer")

def extract_summary(file_path, sheet_name, start_cell=""):
    print(f"DEBUG: file={file_path}, sheet={sheet_name}, cell={start_cell if start_cell else 'auto'}")
    start_cell = start_cell or ""

    # Reproduce the source Excel formatting exactly (fills, fonts, borders,
    # alignment, merged cells, number formats).
    try:
        exact_html = render_excel_range_html(file_path, sheet_name, start_cell)
        if exact_html:
            print(f"DEBUG: exact summary_length={len(exact_html)}")
            return exact_html
    except Exception as e:
        print(f"DEBUG: exact render failed - {e}")

    # Fallback for formats that carry no cell styling (CSV, or .xlsb when exact
    # rendering is disabled): a plain, faithful data table - no invented colours.
    try:
        df = detect_active_range(file_path, sheet_name, start_cell)
        if df is None:
            print("DEBUG: No data detected")
            return ""
        html = _plain_table_html(df)
        print(f"DEBUG: fallback summary_length={len(html)}")
        return html
    except Exception as e:
        print(f"DEBUG: extract failed - {e}")
        return ""


def resolve_summary(file_path, sheet_name, start_cell="", summary_format="table"):
    """Return (html_fragment, inline_images) for the {{Summary}} placeholder.
    inline_images is a list of (cid, png_bytes) for image mode; [] for table
    mode. Image mode falls back to the table if the render fails."""
    if summary_format == "image":
        try:
            png = render_excel_range_image(file_path, sheet_name, start_cell or "")
            if png:
                return (
                    '<img src="cid:summaryimg" alt="Summary" '
                    'style="max-width:100%;height:auto;border:0;display:block;">',
                    [("summaryimg", png)],
                )
        except Exception as e:
            print(f"DEBUG: summary image render failed - {e}")
    return extract_summary(file_path, sheet_name, start_cell), []


def send_email(smtp_config, to_list, subject, html_body, cc_list=None, attachments=None,
               inline_images=None):
    try:
        msg = MIMEMultipart("mixed")
        msg["From"] = f"{smtp_config.get('sender_name', '')} <{smtp_config['sender_email']}>"
        msg["To"] = ", ".join(to_list)
        msg["Subject"] = subject

        if cc_list:
            msg["Cc"] = ", ".join(cc_list)

        if inline_images:
            related = MIMEMultipart("related")
            related.attach(MIMEText(html_body, "html"))
            for cid, img_bytes in inline_images:
                subtype = "gif" if img_bytes[:6] in (b"GIF87a", b"GIF89a") else "png"
                img = MIMEImage(img_bytes, _subtype=subtype)
                img.add_header("Content-ID", f"<{cid}>")
                img.add_header("Content-Disposition", "inline", filename=f"{cid}.{subtype}")
                related.attach(img)
            msg.attach(related)
        else:
            msg.attach(MIMEText(html_body, "html"))

        if attachments:
            for file_path in attachments:
                if os.path.exists(file_path):
                    with open(file_path, "rb") as f:
                        part = MIMEBase("application", "octet-stream")
                        part.set_payload(f.read())
                        encoders.encode_base64(part)
                        part.add_header("Content-Disposition", f'attachment; filename="{os.path.basename(file_path)}"')
                        msg.attach(part)
        
        server = smtplib.SMTP(smtp_config["smtp_server"], smtp_config["smtp_port"], timeout=30)
        server.starttls()
        server.login(smtp_config["sender_email"], smtp_config["password"])
        
        all_recipients = list(to_list)
        if cc_list:
            all_recipients.extend(cc_list)
        if smtp_config.get("sender_email"):
            all_recipients.append(smtp_config["sender_email"])
        
        server.sendmail(smtp_config["sender_email"], all_recipients, msg.as_string())
        server.quit()
        return {"success": True, "message": "Sent"}
    except Exception as e:
        return {"success": False, "message": str(e)}

def _format_duration(seconds: float) -> str:
    seconds = max(0.0, seconds)
    if seconds < 10:
        return f"{seconds:.1f}s"
    if seconds < 60:
        return f"{int(seconds)}s"
    minutes, secs = divmod(int(seconds), 60)
    if minutes < 60:
        return f"{minutes}m {secs}s"
    hours, minutes = divmod(minutes, 60)
    return f"{hours}h {minutes}m"


def _format_when(dt: "datetime") -> str:
    return dt.strftime("%b %d, %Y at %-I:%M %p")


def _smtp_config_from_profile(profile) -> dict:
    return {
        "smtp_server": profile.smtp_server,
        "smtp_port": profile.smtp_port,
        "sender_email": profile.sender_email,
        "sender_name": profile.sender_name,
        "password": profile.password,
    }


_FREQUENCY_LABELS = {"once": "One-time", "daily": "Daily", "weekly": "Weekly", "monthly": "Monthly"}

_EMAIL_FOOTER = """
          <tr><td style="padding:0 32px 24px 32px;">
            <p style="margin:0 0 6px 0;color:#71717a;font-size:12px;text-align:center;">Thank you for using Automation Studio! 🎉</p>
            <p style="margin:0;color:#a1a1aa;font-size:11px;text-align:center;">You can turn this off in Settings.</p>
          </td></tr>
"""


def _build_schedule_confirmation_email(schedule_name: str, frequency: str, next_run) -> str:
    """Blue, calendar-themed card confirming a campaign has been scheduled -
    visually distinct from the completion (black/green) and reminder
    (amber) cards so the three are easy to tell apart at a glance."""
    link = f"{app_settings.FRONTEND_URL}/schedules"
    frequency_label = _FREQUENCY_LABELS.get(frequency, frequency.title())
    return f"""
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#f4f4f7;padding:32px 0;font-family:-apple-system,'Segoe UI',Arial,sans-serif;">
      <tr><td align="center">
        <table role="presentation" width="480" cellpadding="0" cellspacing="0" style="background:#ffffff;border-radius:16px;overflow:hidden;">
          <tr><td style="background:#1d4ed8;padding:24px 32px;">
            <span style="color:#ffffff;font-size:18px;font-weight:700;">📅 Automation Studio</span>
          </td></tr>
          <tr><td align="center" style="padding:20px 0 0 0;">
            <img src="cid:mascot" width="90" height="75" alt="" style="display:block;border:0;">
          </td></tr>
          <tr><td style="padding:12px 32px 4px 32px;">
            <span style="display:inline-block;padding:5px 12px;border-radius:999px;background:#eff6ff;color:#1d4ed8;font-size:12px;font-weight:600;">🗓️ Scheduled</span>
            <h1 style="margin:14px 0 4px 0;font-size:19px;color:#0A0A0A;">{schedule_name}</h1>
            <p style="margin:0;color:#71717a;font-size:13px;">Your campaign has been scheduled and will run automatically.</p>
          </td></tr>
          <tr><td style="padding:20px 32px 8px 32px;">
            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border:1px solid #f0f0f0;border-radius:10px;">
              <tr>
                <td style="padding:12px 16px;font-size:12px;color:#a1a1aa;border-bottom:1px solid #f4f4f5;">Frequency</td>
                <td style="padding:12px 16px;font-size:13px;color:#0A0A0A;font-weight:600;border-bottom:1px solid #f4f4f5;text-align:right;">{frequency_label}</td>
              </tr>
              <tr>
                <td style="padding:12px 16px;font-size:12px;color:#a1a1aa;">Next run</td>
                <td style="padding:12px 16px;font-size:13px;color:#0A0A0A;font-weight:600;text-align:right;">{_format_when(next_run)}</td>
              </tr>
            </table>
          </td></tr>
          <tr><td style="padding:16px 32px 32px 32px;">
            <a href="{link}" style="display:block;text-align:center;background:#1d4ed8;color:#ffffff;text-decoration:none;padding:13px 0;border-radius:10px;font-size:14px;font-weight:600;">Manage Schedules &rarr;</a>
          </td></tr>
{_EMAIL_FOOTER}
        </table>
      </td></tr>
    </table>
    """


def _build_schedule_reminder_email(schedule_name: str, next_run) -> str:
    """Amber, clock-themed "heads up" card sent 15 minutes before a scheduled
    campaign fires."""
    link = f"{app_settings.FRONTEND_URL}/schedules"
    return f"""
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#f4f4f7;padding:32px 0;font-family:-apple-system,'Segoe UI',Arial,sans-serif;">
      <tr><td align="center">
        <table role="presentation" width="480" cellpadding="0" cellspacing="0" style="background:#ffffff;border-radius:16px;overflow:hidden;">
          <tr><td style="background:#b45309;padding:24px 32px;">
            <span style="color:#ffffff;font-size:18px;font-weight:700;">⏰ Automation Studio</span>
          </td></tr>
          <tr><td align="center" style="padding:20px 0 0 0;">
            <img src="cid:mascot" width="90" height="75" alt="" style="display:block;border:0;">
          </td></tr>
          <tr><td style="padding:12px 32px 4px 32px;">
            <span style="display:inline-block;padding:5px 12px;border-radius:999px;background:#fffbeb;color:#b45309;font-size:12px;font-weight:600;">⏳ Starting in 15 minutes</span>
            <h1 style="margin:14px 0 4px 0;font-size:19px;color:#0A0A0A;">{schedule_name}</h1>
            <p style="margin:0;color:#71717a;font-size:13px;">Scheduled to run at {_format_when(next_run)}.</p>
          </td></tr>
          <tr><td style="padding:20px 32px 8px 32px;">
            <p style="margin:0;padding:14px 16px;background:#fffbeb;border-radius:10px;font-size:13px;color:#78350f;">This will send automatically - no action needed. Head to Schedules if you want to reschedule or cancel it first.</p>
          </td></tr>
          <tr><td style="padding:16px 32px 32px 32px;">
            <a href="{link}" style="display:block;text-align:center;background:#0A0A0A;color:#ffffff;text-decoration:none;padding:13px 0;border-radius:10px;font-size:14px;font-weight:600;">View Schedule &rarr;</a>
          </td></tr>
{_EMAIL_FOOTER}
        </table>
      </td></tr>
    </table>
    """


def _build_campaign_summary_email(campaign_name: str, execution, failed: list) -> str:
    """A branded, table-based (Outlook-safe) HTML card: status, how long the
    send took, sent/failed/total counts, the failed-branch table when there
    are any, and a deep link straight back to this execution in History."""
    duration = (
        _format_duration((execution.completed_at - execution.created_at).total_seconds())
        if execution.completed_at and execution.created_at else "—"
    )
    total = execution.total_emails or (execution.sent_count + execution.failed_count)
    ok = len(failed) == 0
    status_label = "Completed successfully" if ok else "Completed with failures"
    status_bg, status_color = ("#ecfdf5", "#059669") if ok else ("#fef2f2", "#dc2626")
    status_icon = "✅" if ok else "⚠️"
    failed_bg, failed_color = ("#f4f4f5", "#71717a") if ok else ("#fef2f2", "#dc2626")
    link = f"{app_settings.FRONTEND_URL}/history?execution={execution.id}"

    failed_section = ""
    if failed:
        rows = "".join(
            f'<tr><td style="padding:8px 10px;border-bottom:1px solid #f0f0f0;font-size:13px;color:#27272a;">{r["branch"]}</td>'
            f'<td style="padding:8px 10px;border-bottom:1px solid #f0f0f0;font-size:12px;color:#dc2626;">{r["reason"]}</td></tr>'
            for r in failed
        )
        failed_section = f"""
          <tr><td style="padding:0 32px 8px 32px;">
            <p style="margin:0 0 8px 0;font-size:13px;font-weight:600;color:#27272a;">Failed branches</p>
            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border:1px solid #f0f0f0;border-radius:8px;overflow:hidden;">
              <tr><th style="padding:8px 10px;text-align:left;font-size:11px;color:#a1a1aa;background:#fafafa;">Branch</th>
                  <th style="padding:8px 10px;text-align:left;font-size:11px;color:#a1a1aa;background:#fafafa;">Reason</th></tr>
              {rows}
            </table>
          </td></tr>
        """

    return f"""
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#f4f4f7;padding:32px 0;font-family:-apple-system,'Segoe UI',Arial,sans-serif;">
      <tr><td align="center">
        <table role="presentation" width="480" cellpadding="0" cellspacing="0" style="background:#ffffff;border-radius:16px;overflow:hidden;">
          <tr><td style="background:#0A0A0A;padding:24px 32px;">
            <span style="color:#ffffff;font-size:18px;font-weight:700;">⚡ Automation Studio</span>
          </td></tr>
          <tr><td align="center" style="padding:20px 0 0 0;">
            <img src="cid:mascot" width="90" height="75" alt="" style="display:block;border:0;">
          </td></tr>
          <tr><td style="padding:12px 32px 4px 32px;">
            <span style="display:inline-block;padding:5px 12px;border-radius:999px;background:{status_bg};color:{status_color};font-size:12px;font-weight:600;">{status_icon} {status_label}</span>
            <h1 style="margin:14px 0 4px 0;font-size:19px;color:#0A0A0A;">{campaign_name}</h1>
            <p style="margin:0;color:#71717a;font-size:13px;">Finished in {duration}</p>
          </td></tr>
          <tr><td style="padding:20px 32px 8px 32px;">
            <table role="presentation" width="100%" cellpadding="0" cellspacing="0">
              <tr>
                <td width="33%" align="center" style="padding:16px 6px;background:#fafafa;border-radius:12px;">
                  <div style="font-size:22px;font-weight:700;color:#0A0A0A;">{total}</div>
                  <div style="font-size:11px;color:#71717a;margin-top:2px;">Total</div>
                </td>
                <td width="3%"></td>
                <td width="33%" align="center" style="padding:16px 6px;background:#ecfdf5;border-radius:12px;">
                  <div style="font-size:22px;font-weight:700;color:#059669;">{execution.sent_count}</div>
                  <div style="font-size:11px;color:#059669;margin-top:2px;">Sent</div>
                </td>
                <td width="3%"></td>
                <td width="33%" align="center" style="padding:16px 6px;background:{failed_bg};border-radius:12px;">
                  <div style="font-size:22px;font-weight:700;color:{failed_color};">{execution.failed_count}</div>
                  <div style="font-size:11px;color:{failed_color};margin-top:2px;">Failed</div>
                </td>
              </tr>
            </table>
          </td></tr>
          {failed_section}
          <tr><td style="padding:16px 32px 32px 32px;">
            <a href="{link}" style="display:block;text-align:center;background:#0A0A0A;color:#ffffff;text-decoration:none;padding:13px 0;border-radius:10px;font-size:14px;font-weight:600;">View in Studio &rarr;</a>
          </td></tr>
{_EMAIL_FOOTER}
        </table>
      </td></tr>
    </table>
    """


def _notify_campaign_result(db: Session, user_id: int, execution, results, profile):
    """Create an in-app notification for the campaign owner, and - if they
    have completion emails enabled (the default) - send a branded summary
    card too, so results don't go unnoticed until someone happens to open
    History. Fires for every completed campaign, not just failed ones."""
    failed = [r for r in results if r["status"] == "failed"]
    campaign_name = execution.campaign_name or "Campaign"

    if failed:
        title = f"{len(failed)} email(s) failed in \"{campaign_name}\""
        message = f"{execution.sent_count} sent, {execution.failed_count} failed."
        notif_type = "campaign_failed"
    else:
        title = f"\"{campaign_name}\" sent successfully"
        message = f"All {execution.sent_count} email(s) sent."
        notif_type = "campaign_completed"

    db.add(Notification(
        user_id=user_id, type=notif_type, title=title, message=message,
        link=f"/history?execution={execution.id}",
    ))
    db.commit()

    setting = db.query(Setting).filter_by(user_id=user_id).first()
    notify_enabled = setting.notify_on_completion if (setting and setting.notify_on_completion is not None) else True
    if not notify_enabled:
        return

    user = db.query(User).filter_by(id=user_id).first()
    if not user or not user.email:
        return

    subject = f"✅ \"{campaign_name}\" completed" if not failed else f"⚠️ \"{campaign_name}\" completed with {len(failed)} failure(s)"
    body = _build_campaign_summary_email(campaign_name, execution, failed)
    mascot = generate_bounce_gif((5, 150, 105), "happy") if not failed else generate_bounce_gif((220, 38, 38), "alert")
    send_email(_smtp_config_from_profile(profile), [user.email], subject, body, inline_images=[("mascot", mascot)])


def _send_branch_email(db: Session, request: CampaignExecuteRequest, profile, execution,
                        branch: str, to_list: list, cc_list: list, existing_log=None) -> dict:
    """Resolve the {{Summary}} block and send one email for one branch. Shared
    by run_campaign() (creates a fresh EmailLog) and the retry endpoint
    (updates an existing failed one in place) so both paths stay identical."""
    file_path = None
    for ext in ['.xlsx', '.xls', '.xlsb', '.csv']:
        test_path = os.path.join(request.campaign_folder, f"{branch}{ext}")
        if os.path.exists(test_path):
            file_path = test_path
            break
    if not file_path:
        file_path = os.path.join(request.campaign_folder, f"{branch}.xlsx")

    summary_html, inline_images = "", []
    if os.path.exists(file_path):
        summary_html, inline_images = resolve_summary(
            file_path, request.sheet_name, request.start_cell, request.summary_format
        )

    subject = request.subject.replace("{{BranchName}}", branch).replace("{{ReportType}}", request.report_type)
    body = request.body_template.replace("{{BranchName}}", branch).replace("{{ReportType}}", request.report_type).replace("{{SenderName}}", profile.sender_name or "").replace("{{Summary}}", summary_html)
    subject = _resolve_cell_placeholders(subject, file_path, request.sheet_name)
    body = _resolve_cell_placeholders(body, file_path, request.sheet_name)
    body = body.replace("\n", "<br>")

    attachments = [file_path] if request.attach_file and os.path.exists(file_path) else []
    result = send_email(_smtp_config_from_profile(profile), to_list, subject, body, cc_list, attachments,
                        inline_images=inline_images)

    now = datetime.utcnow()
    if existing_log is not None:
        existing_log.subject = subject
        existing_log.status = "sent" if result["success"] else "failed"
        existing_log.error_message = result.get("message", "") if not result["success"] else ""
        existing_log.sent_at = now if result["success"] else None
        existing_log.attempted_at = now
    else:
        db.add(EmailLog(
            execution_id=execution.id,
            branch_name=branch,
            recipient_to=", ".join(to_list),
            recipient_cc=", ".join(cc_list),
            subject=subject,
            status="sent" if result["success"] else "failed",
            error_message=result.get("message", "") if not result["success"] else "",
            sent_at=now if result["success"] else None,
            attempted_at=now,
        ))
    db.commit()

    return {
        "branch": branch,
        "status": "sent" if result["success"] else "failed",
        "reason": result.get("message", "") if not result["success"] else "",
    }


def run_campaign(db: Session, user_id: int, request: CampaignExecuteRequest) -> dict:
    """Send one email per mapping entry for this campaign config. Shared by the
    execute-now endpoint and the schedule runner (app/core/scheduler.py)."""
    # Get SMTP profile
    profile = db.query(SMTPProfile).filter_by(profile_name=request.smtp_profile).first()
    if not profile:
        raise HTTPException(status_code=404, detail="SMTP profile not found")
    
    # Get mapping entries
    entries = db.query(MappingEntry).filter_by(mapping_id=request.mapping_id).all() if request.mapping_id else []
    if not entries:
        raise HTTPException(status_code=404, detail="No mapping entries found")
    
    # Create execution record
    execution = Execution(
        user_id=user_id,
        campaign_name=request.report_type,
        campaign_config=request.model_dump_json(),
        status="in_progress",
        send_method="SMTP",
        mode="static/static",
        total_emails=len(entries),
        sent_count=0,
        failed_count=0
    )
    db.add(execution)
    db.commit()

    # Warm the LibreOffice-backed caches for the whole campaign in one call each
    # (soffice startup dominates), so the per-recipient loop just reads them.
    if request.campaign_folder and os.path.isdir(request.campaign_folder):
        folder_files = [
            os.path.join(request.campaign_folder, f)
            for f in os.listdir(request.campaign_folder)
        ]
        xlsb_files = [f for f in folder_files if f.lower().endswith(".xlsb")]
        if xlsb_files:
            prewarm_xlsb_conversions(xlsb_files)
        if request.summary_format == "image":
            sheet_files = [f for f in folder_files
                           if f.lower().endswith((".xlsx", ".xls", ".xlsb"))]
            if sheet_files:
                prewarm_summary_images(sheet_files, request.sheet_name, request.start_cell)

    # Process each entry
    results = []
    for entry in entries:
        branch = entry.branch_name
        # Strip any file extension from branch name
        import re as re_module
        branch = re_module.sub(r'\.(xlsx|xls|xlsb|csv)$', '', branch, flags=re_module.IGNORECASE)
        to_list = [e.strip() for e in entry.to_recipients.replace(";", ",").split(",") if e.strip()]
        cc_raw = str(entry.cc_recipients or "")
        cc_list = [e.strip() for e in cc_raw.replace(";", ",").split(",") if e.strip()] if cc_raw and cc_raw.lower() != "nan" else []

        results.append(_send_branch_email(db, request, profile, execution, branch, to_list, cc_list))
    
    # Update execution
    execution.status = "completed"
    execution.completed_at = datetime.utcnow()
    execution.sent_count = sum(1 for r in results if r["status"] == "sent")
    execution.failed_count = sum(1 for r in results if r["status"] == "failed")
    db.commit()

    try:
        _notify_campaign_result(db, user_id, execution, results, profile)
    except Exception as notify_err:
        print(f"DEBUG: campaign notification failed - {notify_err}")

    return {"success": True, "sent": execution.sent_count, "failed": execution.failed_count, "results": results}


@router.post("/execute-now")
def execute_campaign(request: CampaignExecuteRequest, db: Session = Depends(get_db), auth: tuple = Depends(get_current_user)):
    user_id, role = auth
    return run_campaign(db, user_id, request)


@router.post("/preview-summary")
async def preview_summary(
    file: UploadFile = File(...),
    sheet_name: str = Form("Summary"),
    summary_format: str = Form("table"),
    auth: tuple = Depends(get_current_user),
):
    """Render the {{Summary}} block for a single uploaded file so the campaign
    wizard preview matches what recipients will receive."""
    suffix = os.path.splitext(file.filename or "")[1] or ".xlsx"
    print(f"DEBUG: preview-summary sheet={sheet_name!r} format={summary_format!r} file={file.filename!r}")
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    try:
        tmp.write(await file.read())
        tmp.close()
        if summary_format == "image":
            try:
                png = render_excel_range_image(tmp.name, sheet_name, "")
                print(f"DEBUG: preview image render -> {len(png) if png else 0} bytes")
            except Exception as e:
                print(f"DEBUG: preview image render failed - {e}")
                png = None
            if png:
                b64 = base64.b64encode(png).decode("ascii")
                return {"html": f'<img src="data:image/png;base64,{b64}" '
                                'style="max-width:100%;height:auto;" alt="Summary">'}
        html = extract_summary(tmp.name, sheet_name, "")
        return {"html": html or ""}
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass
