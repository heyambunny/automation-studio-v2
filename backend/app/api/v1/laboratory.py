import os
import io
import re
import zipfile
import tempfile
from copy import copy
from fastapi import APIRouter, Depends, HTTPException, Header, UploadFile, File, Form
from fastapi.responses import StreamingResponse
import openpyxl
import pandas as pd
from app.core.security import decode_token
from app.services.excel_service import _converted_xlsx

router = APIRouter(prefix="/laboratory", tags=["laboratory"])

MAX_RECOMMENDED_GROUPS = 300


def get_current_user(authorization: str = Header(...)):
    token = authorization.replace("Bearer ", "")
    payload = decode_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="Invalid token")
    return int(payload.get("sub")), payload.get("role", "viewer")


def _sanitize_filename(value: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "_", str(value)).strip()
    return (name or "Unnamed")[:100]


def _save_upload(file: UploadFile) -> str:
    suffix = os.path.splitext(file.filename or "")[1].lower() or ".xlsx"
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    tmp.write(file.file.read())
    tmp.close()
    return tmp.name


def _resolve_workbook_path(tmp_path: str) -> str:
    """.xlsb has no Python reader that exposes cell styles, so route it
    through the same LibreOffice conversion the rest of the app already uses
    for exact-formatting rendering."""
    if tmp_path.lower().endswith(".xlsb"):
        converted = _converted_xlsx(tmp_path)
        if not converted:
            raise HTTPException(status_code=400, detail="Could not read this .xlsb file - LibreOffice conversion failed")
        return converted
    return tmp_path


def _find_header_row(ws):
    """First row with any non-empty cell is treated as the header row - master
    reports are expected to be a single flat table per sheet, optionally
    preceded by blank rows."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        if any(c.value not in (None, "") for c in row):
            return row[0].row
    return 1


def _sheet_headers(ws):
    """Returns (header_row_idx, header_cells, last_col, headers)."""
    header_row_idx = _find_header_row(ws)
    header_cells = list(ws[header_row_idx])
    last_col = 0
    for c in header_cells:
        if c.value not in (None, ""):
            last_col = c.column
    headers = [str(ws.cell(row=header_row_idx, column=i).value or "") for i in range(1, last_col + 1)]
    return header_row_idx, header_cells, last_col, headers


def _column_summary(name: str, unique_vals: list, row_count: int) -> dict:
    unique_count = len(unique_vals)
    recommended = 1 < unique_count < max(2, row_count) and unique_count <= MAX_RECOMMENDED_GROUPS
    return {
        "name": name,
        "unique_count": unique_count,
        "sample_values": unique_vals[:8],
        "recommended": recommended,
    }


def _copy_cell_style(src, dest):
    dest.font = copy(src.font)
    dest.fill = copy(src.fill)
    dest.border = copy(src.border)
    dest.alignment = copy(src.alignment)
    dest.number_format = src.number_format


@router.post("/analyze")
async def analyze_file(
    file: UploadFile = File(...),
    auth: tuple = Depends(get_current_user),
):
    """Read every sheet of an uploaded master file and report only the
    columns that exist in ALL of them - those are the only ones that can
    split every sheet consistently, so they're the only ones offered."""
    tmp_path = _save_upload(file)
    try:
        is_csv = tmp_path.lower().endswith(".csv")
        if is_csv:
            df = pd.read_csv(tmp_path, dtype=str)
            headers = [str(c) for c in df.columns]
            row_count = len(df)
            columns = []
            for col in headers:
                values = df[col].dropna().astype(str)
                values = values[values.str.strip() != ""]
                unique_vals = values.unique().tolist()
                columns.append(_column_summary(col, unique_vals, row_count))
            return {"sheet_names": None, "sheets_analyzed": 1, "row_count": row_count, "columns": columns}

        work_path = _resolve_workbook_path(tmp_path)
        wb = openpyxl.load_workbook(work_path, data_only=True)
        sheet_names = wb.sheetnames

        per_sheet = {}
        header_sets = []
        first_sheet_headers = None
        total_rows = 0
        for sname in sheet_names:
            ws = wb[sname]
            header_row_idx, header_cells, last_col, headers = _sheet_headers(ws)
            data_rows = list(ws.iter_rows(min_row=header_row_idx + 1, max_col=last_col, values_only=True)) if last_col else []
            per_sheet[sname] = {"headers": headers, "data_rows": data_rows}
            header_sets.append(set(headers))
            if first_sheet_headers is None:
                first_sheet_headers = headers
            total_rows += len(data_rows)

        common_columns = set.intersection(*header_sets) if header_sets else set()
        # keep the first sheet's column order for anything that's common to all
        ordered_common = [c for c in (first_sheet_headers or []) if c in common_columns and c]

        columns = []
        for col_name in ordered_common:
            all_values: list = []
            for sname in sheet_names:
                headers = per_sheet[sname]["headers"]
                idx = headers.index(col_name)
                for r in per_sheet[sname]["data_rows"]:
                    v = r[idx]
                    if v not in (None, "") and str(v).strip() != "":
                        all_values.append(str(v).strip())
            unique_vals = list(dict.fromkeys(all_values))
            columns.append(_column_summary(col_name, unique_vals, total_rows))

        return {"sheet_names": sheet_names, "sheets_analyzed": len(sheet_names), "row_count": total_rows, "columns": columns}
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


@router.post("/split")
async def split_file(
    file: UploadFile = File(...),
    column_name: str = Form(...),
    auth: tuple = Depends(get_current_user),
):
    """Split every sheet of the uploaded file by column_name at once: each
    output file gets one branch's rows from every original sheet, so a
    multi-sheet master report (e.g. Summary/Enrollment/Redemption) produces
    one multi-sheet file per branch instead of splitting only one sheet."""
    tmp_path = _save_upload(file)
    try:
        is_csv = tmp_path.lower().endswith(".csv")
        zip_buffer = io.BytesIO()

        if is_csv:
            df = pd.read_csv(tmp_path, dtype=str)
            if column_name not in df.columns:
                raise HTTPException(status_code=400, detail=f"Column \"{column_name}\" not found")
            groups = df.groupby(df[column_name].fillna("Unassigned").astype(str).str.strip())
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                for group_value, group_df in groups:
                    out = io.StringIO()
                    group_df.to_csv(out, index=False)
                    zf.writestr(f"{_sanitize_filename(group_value)}.csv", out.getvalue())
        else:
            work_path = _resolve_workbook_path(tmp_path)
            wb = openpyxl.load_workbook(work_path, data_only=True)
            sheet_names = wb.sheetnames

            sheet_infos = {}
            for sname in sheet_names:
                ws = wb[sname]
                header_row_idx, header_cells, last_col, headers = _sheet_headers(ws)
                if column_name not in headers:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Column \"{column_name}\" is missing from sheet \"{sname}\" - it must exist in every sheet to split by it",
                    )
                col_idx = headers.index(column_name) + 1
                col_widths = {
                    i: ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width
                    for i in range(1, last_col + 1) if openpyxl.utils.get_column_letter(i) in ws.column_dimensions
                }
                sheet_infos[sname] = {
                    "header_cells": header_cells,
                    "last_col": last_col,
                    "col_widths": col_widths,
                    "rows_by_group": {},
                }
                for row in ws.iter_rows(min_row=header_row_idx + 1, max_col=last_col):
                    key_cell = row[col_idx - 1]
                    if key_cell.value in (None, "") or str(key_cell.value).strip() == "":
                        continue
                    key = str(key_cell.value).strip()
                    sheet_infos[sname]["rows_by_group"].setdefault(key, []).append(row)

            group_order: list = []
            seen = set()
            for sname in sheet_names:
                for key in sheet_infos[sname]["rows_by_group"]:
                    if key not in seen:
                        seen.add(key)
                        group_order.append(key)

            if not group_order:
                raise HTTPException(status_code=400, detail=f"No data rows found under column \"{column_name}\" in any sheet")

            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                for group_value in group_order:
                    out_wb = openpyxl.Workbook()
                    out_wb.remove(out_wb.active)

                    for sname in sheet_names:
                        info = sheet_infos[sname]
                        out_ws = out_wb.create_sheet(title=(sname or "Sheet1")[:31])

                        for i in range(1, info["last_col"] + 1):
                            letter = openpyxl.utils.get_column_letter(i)
                            if i in info["col_widths"]:
                                out_ws.column_dimensions[letter].width = info["col_widths"][i]

                        for c in info["header_cells"][:info["last_col"]]:
                            dest = out_ws.cell(row=1, column=c.column, value=c.value)
                            _copy_cell_style(c, dest)

                        rows = info["rows_by_group"].get(group_value, [])
                        for r_idx, row in enumerate(rows, start=2):
                            for c in row:
                                dest = out_ws.cell(row=r_idx, column=c.column, value=c.value)
                                _copy_cell_style(c, dest)

                    file_bytes = io.BytesIO()
                    out_wb.save(file_bytes)
                    zf.writestr(f"{_sanitize_filename(group_value)}.xlsx", file_bytes.getvalue())

        zip_buffer.seek(0)
        return StreamingResponse(
            zip_buffer,
            media_type="application/zip",
            headers={"Content-Disposition": "attachment; filename=split_files.zip"},
        )
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
