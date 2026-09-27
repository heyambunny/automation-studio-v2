"""Kitchen: learns a reusable "recipe" from one or more finished example
branch reports (a raw-data sheet + a Summary sheet full of live formulas
referencing it), then re-applies that recipe to every other branch's raw
data to auto-generate their Summary sheets.

Scope (confirmed with the user): a fixed list of aggregate functions
(SUM/SUMIF/SUMIFS/COUNTIF/COUNTIFS/COUNTA/AVERAGE/AVERAGEIF/AVERAGEIFS/MAX/MIN)
referencing the raw-data sheet are fully generalized. Formulas that only
reference other Summary cells (no raw-data crossing) are copied as-is, since
the Summary layout is identical across branches. Anything else (VLOOKUP,
INDEX/MATCH, nested IF, array formulas, or a shape that disagrees between
training examples) is flagged for manual review rather than guessed at.
"""
import re
from copy import copy, deepcopy

import openpyxl

SUPPORTED_AGG_FUNCTIONS = {
    "SUM", "SUMIF", "SUMIFS", "COUNTIF", "COUNTIFS", "COUNTA",
    "AVERAGE", "AVERAGEIF", "AVERAGEIFS", "MAX", "MIN",
}

# Which argument positions are "a range to aggregate" vs "a criteria range"
# vs "a criteria value", per function - needed for both classification and
# for writing a readable plain-English description. SUMIFS/AVERAGEIFS/COUNTIFS
# repeat (criteria_range, criteria) pairs after their fixed leading args.
_FIXED_ARG_ROLES = {
    "SUM": ["range"],
    "COUNTA": ["range"],
    "AVERAGE": ["range"],
    "MAX": ["range"],
    "MIN": ["range"],
    "SUMIF": ["criteria_range", "criteria", "range"],
    "AVERAGEIF": ["criteria_range", "criteria", "range"],
    "COUNTIF": ["criteria_range", "criteria"],
}
_PAIRED_FUNCTIONS = {
    "SUMIFS": "range",       # SUMIFS(sum_range, crit_range, crit, ...)
    "AVERAGEIFS": "range",   # AVERAGEIFS(avg_range, crit_range, crit, ...)
    "COUNTIFS": None,        # COUNTIFS(crit_range, crit, ...) - no leading range
}

_RANGE_TOKEN_RE = re.compile(
    r"^(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_ ]*))!(\$?[A-Z]{1,3}\$?\d*)(?::(\$?[A-Z]{1,3}\$?\d*))?$"
)
_COL_LETTER_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d*)$")


def _split_top_level_args(s: str) -> list:
    """Split a function's argument string on top-level commas, ignoring
    commas inside nested parens or double-quoted strings."""
    args, depth, in_quotes, current = [], 0, False, []
    for ch in s:
        if ch == '"':
            in_quotes = not in_quotes
            current.append(ch)
        elif in_quotes:
            current.append(ch)
        elif ch == "(":
            depth += 1
            current.append(ch)
        elif ch == ")":
            depth -= 1
            current.append(ch)
        elif ch == "," and depth == 0:
            args.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
    if current:
        args.append("".join(current).strip())
    return args


def parse_formula(formula):
    """`=FUNC(arg1,arg2,...)` -> (FUNC, [arg1, arg2, ...]); None if it isn't
    a single top-level function call (covers plain arithmetic like
    `=B5/B6*100`, which is handled separately as an "internal" formula)."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    body = formula[1:].strip()
    m = re.match(r"^([A-Za-z]+)\((.*)\)$", body, re.DOTALL)
    if not m:
        return None
    depth = 0
    for i, ch in enumerate(body):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
            if depth == 0 and i != len(body) - 1:
                return None  # trailing content after the outer call
    return m.group(1).upper(), _split_top_level_args(m.group(2))


def parse_range_ref(token: str):
    """"'Raw Data'!C2:C450" or "RawData!C:C" -> {"sheet","start","end"}, or
    None if this argument isn't a sheet-qualified range at all (e.g. a
    literal or a plain criteria value)."""
    m = _RANGE_TOKEN_RE.match(token.strip())
    if not m:
        return None
    return {
        "sheet": m.group(1) or m.group(2),
        "start": m.group(3),
        "end": m.group(4) or m.group(3),
    }


def _col_letter_of(ref: str):
    m = _COL_LETTER_RE.match(ref.strip())
    return m.group(1) if m else None


def _find_header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        if any(c.value not in (None, "") for c in row):
            return row[0].row
    return 1


def sheet_headers(ws):
    """Returns (header_row_idx, last_col, headers) - same shape as
    laboratory.py's helper of the same name, duplicated here so this module
    has no import dependency on the API layer (laboratory.py imports this
    module, not the other way around)."""
    header_row_idx = _find_header_row(ws)
    last_col = 0
    for c in ws[header_row_idx]:
        if c.value not in (None, ""):
            last_col = c.column
    headers = [str(ws.cell(row=header_row_idx, column=i).value or "") for i in range(1, last_col + 1)]
    return header_row_idx, last_col, headers


def _header_for_column(headers, col_letter):
    if not col_letter:
        return None
    idx = openpyxl.utils.column_index_from_string(col_letter) - 1
    return headers[idx] if 0 <= idx < len(headers) and headers[idx] else None


def _column_letter_for_header(headers, header_name):
    if header_name in headers:
        return openpyxl.utils.get_column_letter(headers.index(header_name) + 1)
    return None


def _arg_roles(func_name: str, n_args: int) -> list:
    if func_name in _PAIRED_FUNCTIONS:
        leading = _PAIRED_FUNCTIONS[func_name]
        roles = [leading] if leading else []
        pair_count = (n_args - len(roles)) // 2
        for _ in range(pair_count):
            roles += ["criteria_range", "criteria"]
        return roles
    return _FIXED_ARG_ROLES.get(func_name, [])


def _references_another_sheet(formula) -> bool:
    """True if the formula is sheet-qualified at all (contains a `!`, ignoring
    a bare `!=` which Excel formulas never produce) - i.e. it isn't safe to
    copy verbatim into another branch's file, because that other sheet either
    won't exist there or won't hold the same data."""
    return isinstance(formula, str) and formula.startswith("=") and "!" in formula


def _classify_literal_values(values_by_example: dict, branch_by_example: dict) -> dict:
    """values_by_example: {example_id: value}, already comparable (unquoted
    strings, numbers, etc). Identical everywhere -> a true constant. Differs
    but matches each example's own branch name -> that example's branch name
    was hardcoded in. Differs without a clean match -> ambiguous."""
    unique_vals = set(values_by_example.values())
    if len(unique_vals) == 1:
        return {"mode": "constant", "value": next(iter(unique_vals))}
    if all(values_by_example[ex] == branch_by_example.get(ex) for ex in values_by_example):
        return {"mode": "this_branch_name"}
    return {"mode": "low_confidence", "examples": dict(values_by_example)}


def classify_cell(coord, formulas_by_example, raw_headers_by_example, branch_by_example):
    """formulas_by_example: {example_id: raw cell value/formula at this coord}
    (already read with data_only=False). Returns one rule dict."""
    parsed = {ex: parse_formula(f) for ex, f in formulas_by_example.items()}
    first_val = next(iter(formulas_by_example.values()))

    def _classify_as_internal(reason_prefix: str):
        # Not an aggregate we generalize - safe to copy verbatim ONLY if it's
        # identical across every training example. A cell that differs per
        # example but matches that example's own branch name is a hardcoded
        # branch-name label (e.g. a "Kanpur" typed straight into the Summary,
        # not derived from a formula) - freezing the first example's value
        # there would silently mislabel every other branch's report.
        classification = _classify_literal_values(formulas_by_example, branch_by_example)
        if classification["mode"] == "constant":
            return {"coord": coord, "kind": "internal", "value": classification["value"]}
        if classification["mode"] == "this_branch_name":
            return {"coord": coord, "kind": "internal_branch_name"}
        return {
            "coord": coord, "kind": "needs_review", "example_formula": first_val,
            "reason": f"{reason_prefix} differs between examples without matching each branch's name",
        }

    if any(p is None for p in parsed.values()):
        # Not a single top-level function call in at least one example - if it
        # references ANY sheet (the raw sheet, or something else entirely like
        # a rate/target table) it's a shape we don't understand; otherwise it's
        # arithmetic between Summary cells or a plain label/value.
        if any(_references_another_sheet(f) for f in formulas_by_example.values()):
            return {"coord": coord, "kind": "needs_review", "example_formula": first_val, "reason": "not a simple supported function call"}
        return _classify_as_internal("value")

    func_names = {fn for fn, _ in parsed.values()}
    if len(func_names) != 1:
        return {"coord": coord, "kind": "needs_review", "example_formula": first_val, "reason": "function differs between examples"}
    func_name = next(iter(func_names))

    if func_name not in SUPPORTED_AGG_FUNCTIONS:
        if any(_references_another_sheet(f) for f in formulas_by_example.values()):
            return {"coord": coord, "kind": "needs_review", "example_formula": first_val, "reason": f"'{func_name}' isn't a supported aggregate function"}
        return _classify_as_internal(f"'{func_name}' formula")

    arg_lists = {ex: args for ex, (_, args) in parsed.items()}
    if len({len(a) for a in arg_lists.values()}) != 1:
        return {"coord": coord, "kind": "needs_review", "example_formula": first_val, "reason": "argument count differs between examples"}
    n_args = len(next(iter(arg_lists.values())))
    roles = _arg_roles(func_name, n_args)

    arg_specs = []
    for i in range(n_args):
        role = roles[i] if i < len(roles) else "other"
        values = {ex: arg_lists[ex][i].strip() for ex in arg_lists}

        if role in ("range", "criteria_range"):
            range_refs = {ex: parse_range_ref(v) for ex, v in values.items()}
            if any(r is None for r in range_refs.values()):
                arg_specs.append({"role": role, "type": "unresolved", "reason": "not a sheet-qualified range in every example"})
                continue
            headers = {
                ex: _header_for_column(raw_headers_by_example[ex], _col_letter_of(range_refs[ex]["start"]))
                for ex in values
            }
            resolved = set(headers.values())
            if len(resolved) == 1 and next(iter(resolved)) is not None:
                arg_specs.append({"role": role, "type": "range", "column_header": next(iter(resolved))})
            else:
                arg_specs.append({"role": role, "type": "unresolved", "reason": "column header didn't match across examples"})

        elif role == "criteria":
            if all(v.startswith('"') and v.endswith('"') for v in values.values()):
                literals = {ex: v[1:-1] for ex, v in values.items()}
                arg_specs.append({"role": role, "type": "literal", **_classify_literal_values(literals, branch_by_example)})
            else:
                arg_specs.append({"role": role, "type": "unresolved", "reason": "criteria isn't a plain literal in every example"})
        else:
            arg_specs.append({"role": role, "type": "unresolved", "reason": "unrecognized argument shape"})

    if any(a["type"] == "unresolved" for a in arg_specs):
        return {"coord": coord, "kind": "needs_review", "example_formula": first_val, "reason": "couldn't generalize one or more arguments", "partial_args": arg_specs}

    return {"coord": coord, "kind": "aggregate", "function": func_name, "args": arg_specs}


def build_recipe(examples: list, summary_sheet: str, raw_sheet: str) -> dict:
    """examples: [{"id": str, "branch": str, "path": str}, ...] (2+ items).
    Loads each example twice (data_only=False for formulas, data_only=True
    for a sanity-check cached value) and classifies every populated Summary
    cell across all of them at once."""
    formulas_by_example, values_by_example, headers_by_example, branch_by_example = {}, {}, {}, {}
    all_coords = set()

    for ex in examples:
        wb_f = openpyxl.load_workbook(ex["path"], data_only=False)
        wb_v = openpyxl.load_workbook(ex["path"], data_only=True)
        if summary_sheet not in wb_f.sheetnames or raw_sheet not in wb_f.sheetnames:
            raise ValueError(f"Example '{ex['id']}' is missing a '{summary_sheet}' or '{raw_sheet}' sheet")

        ws_f, ws_v = wb_f[summary_sheet], wb_v[summary_sheet]
        formulas, values = {}, {}
        for row in ws_f.iter_rows():
            for cell in row:
                if cell.value is not None:
                    formulas[cell.coordinate] = cell.value
        for row in ws_v.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values[cell.coordinate] = cell.value

        _, _, raw_headers = sheet_headers(wb_f[raw_sheet])

        formulas_by_example[ex["id"]] = formulas
        values_by_example[ex["id"]] = values
        headers_by_example[ex["id"]] = raw_headers
        branch_by_example[ex["id"]] = ex["branch"]
        all_coords |= set(formulas.keys())

    cells = []
    for coord in sorted(all_coords, key=lambda c: (int(re.search(r"\d+", c).group()), c)):
        per_example = {ex: formulas_by_example[ex].get(coord) for ex in formulas_by_example}
        if any(v is None for v in per_example.values()):
            cells.append({"coord": coord, "kind": "needs_review", "reason": "cell is populated in some examples but not others"})
            continue
        rule = classify_cell(coord, per_example, headers_by_example, branch_by_example)
        if rule["kind"] == "needs_review" and "example_value" not in rule:
            first_ex = next(iter(values_by_example))
            rule["example_value"] = values_by_example[first_ex].get(coord)
        cells.append(rule)

    return {
        "summary_sheet": summary_sheet,
        "raw_sheet": raw_sheet,
        "style_source_id": examples[0]["id"],
        "style_source_path": examples[0]["path"],
        "cells": cells,
    }


_FUNCTION_LABEL = {
    "SUM": "Sum", "AVERAGE": "Average", "MAX": "Maximum", "MIN": "Minimum", "COUNTA": "Count of non-blank values in",
    "SUMIF": "Sum", "SUMIFS": "Sum", "AVERAGEIF": "Average", "AVERAGEIFS": "Average",
    "COUNTIF": "Count of rows where", "COUNTIFS": "Count of rows where",
}


def _describe_criteria(arg: dict) -> str:
    if arg["mode"] == "constant":
        return f'"{arg["value"]}"'
    if arg["mode"] == "this_branch_name":
        return "this branch's own name"
    return f"⚠ unclear - varies per branch without a clear pattern ({arg['examples']})"


def describe_rule(rule: dict) -> str:
    coord = rule["coord"]
    if rule["kind"] == "internal":
        return f"{coord}: kept exactly as in the example ({rule.get('value')})"
    if rule["kind"] == "internal_branch_name":
        return f"{coord}: this branch's own name"
    if rule["kind"] == "needs_review":
        return f"{coord}: needs your input - {rule.get('reason', 'unsupported formula')} (was: {rule.get('example_formula', rule.get('example_value', ''))})"

    func = rule["function"]
    args = rule["args"]
    ranges = [a for a in args if a["role"] in ("range",)]
    crit_ranges = [a for a in args if a["role"] == "criteria_range"]
    criteria = [a for a in args if a["role"] == "criteria"]

    if func in ("SUM", "AVERAGE", "MAX", "MIN", "COUNTA"):
        col = ranges[0]["column_header"] if ranges else "?"
        return f'{coord}: {_FUNCTION_LABEL[func]} of "{col}" across all rows'

    label = _FUNCTION_LABEL[func]
    target = f' of "{ranges[0]["column_header"]}"' if ranges and func not in ("COUNTIF", "COUNTIFS") else ""
    conditions = " and ".join(
        f'"{cr["column_header"]}" = {_describe_criteria(cv)}' for cr, cv in zip(crit_ranges, criteria)
    )
    return f"{coord}: {label}{target} where {conditions}"


def is_resolved(rule: dict) -> bool:
    """False for a needs_review cell, or an aggregate cell whose criteria
    couldn't be classified confidently - both need a manual answer (via
    apply_overrides) before generation should be allowed to proceed."""
    if rule["kind"] != "aggregate":
        return rule["kind"] != "needs_review"
    return all(a.get("mode") != "low_confidence" for a in rule["args"] if a["type"] == "literal")


def render_rule_list(recipe: dict) -> list:
    return [
        {"coord": r["coord"], "kind": r["kind"], "description": describe_rule(r), "resolved": is_resolved(r), "rule": r}
        for r in recipe["cells"]
    ]


def apply_overrides(recipe: dict, overrides: dict) -> dict:
    """overrides: {coord: patch}. For a needs_review cell, patch is
    {"kind": "manual_formula", "value": "=..."} - the user's own formula/value,
    copied verbatim into every branch's Summary (same as an "internal" rule).
    For an aggregate cell with a low_confidence criteria arg, patch is
    {"arg_index": i, "mode": "constant"|"this_branch_name", "value": "..." (only for constant)}.
    Returns a new recipe; the input is left untouched."""
    recipe = deepcopy(recipe)
    for i, rule in enumerate(recipe["cells"]):
        patch = overrides.get(rule["coord"])
        if not patch:
            continue
        if rule["kind"] == "needs_review" and patch.get("kind") == "manual_formula":
            recipe["cells"][i] = {"coord": rule["coord"], "kind": "internal", "value": patch["value"]}
        elif rule["kind"] == "aggregate" and "arg_index" in patch:
            arg = rule["args"][patch["arg_index"]]
            arg["mode"] = patch["mode"]
            if patch["mode"] == "constant":
                arg["value"] = patch["value"]
                arg.pop("examples", None)
            elif patch["mode"] == "this_branch_name":
                arg.pop("value", None)
                arg.pop("examples", None)
    return recipe


def regenerate_summary(recipe: dict, style_source_path: str, target_raw_ws, target_headers: list, target_branch_name: str, out_ws):
    """Writes a Summary sheet into out_ws: layout/style copied from the
    recipe's training example, values/formulas rebuilt per-cell from the
    recipe against target_raw_ws (an already-populated branch worksheet, e.g.
    from build_branch_workbooks). Returns a list of coords that fell back to
    a needs_review value (for a visible warning in the response)."""
    style_wb = openpyxl.load_workbook(style_source_path, data_only=False)
    style_ws = style_wb[recipe["summary_sheet"]]

    for row in style_ws.iter_rows():
        for cell in row:
            dest = out_ws.cell(row=cell.row, column=cell.column)
            dest.font, dest.fill, dest.border = copy(cell.font), copy(cell.fill), copy(cell.border)
            dest.alignment, dest.number_format = copy(cell.alignment), cell.number_format
    for col_letter, dim in style_ws.column_dimensions.items():
        if dim.width:
            out_ws.column_dimensions[col_letter].width = dim.width
    for merged_range in style_ws.merged_cells.ranges:
        out_ws.merge_cells(str(merged_range))

    raw_sheet_name = recipe["raw_sheet"]
    target_last_row = target_raw_ws.max_row
    needs_review_coords = []

    for rule in recipe["cells"]:
        coord = rule["coord"]
        if rule["kind"] == "internal":
            out_ws[coord] = rule["value"]
        elif rule["kind"] == "internal_branch_name":
            out_ws[coord] = target_branch_name
        elif rule["kind"] == "needs_review":
            out_ws[coord] = rule.get("example_value", rule.get("example_formula"))
            needs_review_coords.append(coord)
        elif rule["kind"] == "aggregate":
            parts = []
            for arg in rule["args"]:
                if arg["type"] == "range":
                    col_letter = _column_letter_for_header(target_headers, arg["column_header"])
                    if not col_letter:
                        needs_review_coords.append(coord)
                        parts = None
                        break
                    parts.append(f"'{raw_sheet_name}'!{col_letter}2:{col_letter}{target_last_row}")
                elif arg["mode"] == "constant":
                    parts.append(f'"{arg["value"]}"')
                elif arg["mode"] == "this_branch_name":
                    parts.append(f'"{target_branch_name}"')
                else:
                    needs_review_coords.append(coord)
                    parts = None
                    break
            if parts is None:
                out_ws[coord] = None
            else:
                out_ws[coord] = f"={rule['function']}({','.join(parts)})"

    return needs_review_coords
